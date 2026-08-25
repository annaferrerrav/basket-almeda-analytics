# -*- coding: utf-8 -*-
"""
Created on Tue Oct 21 17:35:29 2025

@author: AFERRERR
"""
# -*- coding: utf-8 -*-
"""
Created on Thu Oct 16 16:33:52 2025
@author: AFERRERR
"""

import pandas as pd
import glob
import os

# Ruta a la carpeta
path = r"C:\Users\aferrerr\OneDrive - Hospital Clínic de Barcelona\Escritorio\LF2 - 2025-56\LF2\boxscore_jornada"

# Variable del equipo
equip = "BASKET ALMEDA"

# Buscar todos los archivos Excel (xls, xlsx, xlsm, etc.)
excel_files = glob.glob(os.path.join(path, "*.xls*"))

# Lista para guardar los DataFrames
dfs = []

for file in excel_files:
    try:
        df = pd.read_excel(file)
        if "Equip" in df.columns:
            df_filtrado = df[df["Equip"] == equip]
            if not df_filtrado.empty:
                df_filtrado["Archivo"] = os.path.basename(file)
                dfs.append(df_filtrado)
    except Exception as e:
        print(f"⚠️ Error leyendo {file}: {e}")

if dfs:
    resultado = pd.concat(dfs, ignore_index=True)
    print("✅ Datos concatenados del equipo.\n")

    if "Jugador" in resultado.columns:
        # --- Columnas a promediar ---
        columnas_media = [
            "MIN", "PT", "T2C", "T2I", "T3C", "T3I", "TCC", "TCI", "TLC", "TLI",
            "RO", "RD", "RT", "AS", "BR", "BP", "TapF", "TapC", "MT", "FC",
            "FR", "VA", "+/-", "Plays", "Plays/min"
        ]
        columnas_media = [c for c in columnas_media if c in resultado.columns]

        # Contamos partidos jugados
        #partidos_jugados = resultado.groupby("Jugador").size().reset_index(name="Partidos jugados")
        partidos_jugados = (
            resultado.groupby("Jugador")["Archivo"]
            .nunique()
            .reset_index(name="Partidos jugados")
        )
                
        print(partidos_jugados)
        # Si existe la columna "Dorsal", la mantenemos en el resumen
        if "Dorsal" in resultado.columns:
            dorsales = resultado[["Jugador", "Dorsal"]].drop_duplicates(subset="Jugador")
        else:
            dorsales = pd.DataFrame(columns=["Jugador", "Dorsal"])

        # Calculamos las medias
        medias = resultado.groupby("Jugador")[columnas_media].mean().reset_index()

        # Unimos datos
        resumen = pd.merge(partidos_jugados, medias, on="Jugador")
        resumen = pd.merge(dorsales, resumen, on="Jugador", how="left")

        # Reordenamos columnas
        cols = ["Jugador", "Dorsal", "Partidos jugados"] + [c for c in columnas_media if c in resumen.columns]
        resumen = resumen[cols]

        # Mostramos resultado
        print("📊 Promedios por jugadora (con Dorsal y partidos jugados):\n")
        print(resumen.round(2))

        # (Opcional) Guardar en Excel
        # resumen.to_excel("promedios_jugadoras.xlsx", index=False)
    else:
        print("⚠️ No se encontró la columna 'Jugador' en los archivos.")
else:
    print("No se encontraron datos para el equipo especificado.")

import numpy as np

# --- Afegim percentatges (si les columnes existeixen) ---
def pct(nume, deno):
    return np.round(np.where(deno > 0, (nume / deno) * 100, 0.0), 4)

if {"T2C","T2I"}.issubset(resumen.columns):
    resumen["T2 (%)"] = pct(resumen["T2C"], resumen["T2I"])

if {"T3C","T3I"}.issubset(resumen.columns):
    resumen["T3 (%)"] = pct(resumen["T3C"], resumen["T3I"])

if {"TLC","TLI"}.issubset(resumen.columns):
    resumen["TL (%)"] = pct(resumen["TLC"], resumen["TLI"])

if {"TCC","TCI"}.issubset(resumen.columns):
    resumen["TC (%)"] = pct(resumen["TCC"], resumen["TCI"])

# Arrodonim una mica
for col in ["T2 (%)","T3 (%)","TL (%)","TC (%)"]:
    if col in resumen.columns:
        resumen[col] = resumen[col].round(2)

# --- Reordenem columnes: % just després del que li toca ---
ordre_desitjat = [
    "Jugador", "Dorsal", "Partidos jugados",
    "MIN", "PT",
    "T2C", "T2I", "T2 (%)",
    "T3C", "T3I", "T3 (%)",
    "TCC", "TCI", "TC (%)",
    "TLC", "TLI", "TL (%)",
    "RO", "RD", "RT", "AS", "BR", "BP", "TapF", "TapC", "MT", "FC", "FR",
    "VA", "+/-", "Plays", "Plays/min"
]

# Ens quedem només amb les que existeixen i després afegim la resta (si n'hi ha)
ordre_final = [c for c in ordre_desitjat if c in resumen.columns] + \
              [c for c in resumen.columns if c not in ordre_desitjat]

resumen = resumen[ordre_final]

import numpy as np

# --- Helper per obtenir el denominador "Total" d'una columna d'intents ---
def total_denominator(df, col_intents, jugador_total="Total"):
    if col_intents not in df.columns:
        return None
    # Busca la fila "Total"
    mask_total = df["Jugador"].astype(str).str.strip().str.casefold() == jugador_total.casefold()
    if mask_total.any():
        val = df.loc[mask_total, col_intents].iloc[0]
        return float(val) if pd.notna(val) else None
    # Si no hi ha fila "Total", fem servir la suma de la resta de jugadores
    return float(df.loc[~mask_total, col_intents].sum())

# --- Calcula %US per a cada tipus de tir ---
pairs = [
    ("T2I",  "% US T2"),
    ("T3I",  "% US T3"),
    ("TCI",  "% US TC"),
    ("TLI",  "% US TLL"),
]

for col_intents, col_pctus in pairs:
    den = total_denominator(resumen, col_intents)
    if den is not None and den > 0 and col_intents in resumen.columns:
        resumen[col_pctus] = (resumen[col_intents] / den) * 100.0
    elif col_intents in resumen.columns:
        # Si no hi ha denominador vàlid, deixem 0.0 (pots canviar a np.nan si prefereixes)
        resumen[col_pctus] = 0.0

# Opcional: per a la fila "Total", posa 100% si el denominador és >0
if "Jugador" in resumen.columns:
    mask_total = resumen["Jugador"].astype(str).str.strip().str.casefold() == "total"
    for _, col_pctus in pairs:
        if col_pctus in resumen.columns:
            # 100% només si hi havia denominador > 0
            origen = [p[0] for p in pairs if p[1] == col_pctus][0]
            den = total_denominator(resumen, origen)
            if den is not None and den > 0:
                resumen.loc[mask_total, col_pctus] = 100.0

# Arrodoneix %US
for _, col_pctus in pairs:
    if col_pctus in resumen.columns:
        resumen[col_pctus] = resumen[col_pctus].round(2)

# --- Reordena: %US just després de cada bloc ---
ordre_desitjat = [
    "Jugador", "Dorsal", "Partidos jugados",
    "MIN", "PT",
    "T2C", "T2I", "T2 (%)", "% US T2",
    "T3C", "T3I", "T3 (%)", "% US T3",
    "TCC", "TCI", "TC (%)", "% US TC",
    "TLC", "TLI", "TL (%)", "% US TLL",
    "RO", "RD", "RT", "AS", "BR", "BP", "TapF", "TapC", "MT", "FC", "FR",
    "VA", "+/-", "Plays", "Plays/min"
]

ordre_final = [c for c in ordre_desitjat if c in resumen.columns] + \
              [c for c in resumen.columns if c not in ordre_desitjat]
resumen = resumen[ordre_final]


from openpyxl import load_workbook
import pandas as pd
import numpy as np
from datetime import datetime
from pathlib import Path
import os, shutil, tempfile

def normalize_text(s):
    return "" if pd.isna(s) else str(s).strip().casefold()

# --- Helper per limitar decimals i formatar percentatges ---
def to_excel_percent(val):
    """Retorna valors en fracció 0..1 per Excel amb 2 decimals"""
    try:
        v = float(val)
    except Exception:
        return None if val in [None, "", "nan"] else val
    if np.isnan(v):
        return None
    if abs(v) > 1.0:   # si ve com 25.56 → 0.2556
        v = v / 100.0
    return round(v, 2)  # màxim 2 decimals

def to_excel_number(val):
    """Qualsevol altre número limitat a 2 decimals"""
    try:
        v = float(val)
    except Exception:
        return None
    if np.isnan(v):
        return None
    return round(v, 2)

# --- Totals equip ---
def team_totals_from_resumen(resumen: pd.DataFrame):
    df = resumen.copy()
    mask_total = df["Jugador"].astype(str).str.strip().str.casefold() == "total" if "Jugador" in df.columns else pd.Series(False, index=df.index)
    if mask_total.any():
        return df.loc[mask_total].iloc[0]

    cols_sum = ["PT","T2C","T2I","T3C","T3I","TCC","TCI","TLC","TLI",
                "RO","RD","RT","AS","BR","BP","TapF","TapC","MT","FC","FR","VA","Plays"]
    totals = {}
    for c in cols_sum:
        if c in df.columns:
            totals[c] = pd.to_numeric(df[c], errors="coerce").fillna(0).sum()

    def pct(n,d):
        return np.round(np.where(d > 0, n / d, 0.0), 4)

    totals["T2 (%)"] = pct(totals.get("T2C",0), totals.get("T2I",0))
    totals["T3 (%)"] = pct(totals.get("T3C",0), totals.get("T3I",0))
    totals["TC (%)"] = pct(totals.get("TCC",0), totals.get("TCI",0))
    totals["TL (%)"] = pct(totals.get("TLC",0), totals.get("TLI",0))
    return pd.Series(totals)

def find_player_blocks(ws, max_search_rows=400):
    header_rows = []
    for r in range(1, max_search_rows+1):
        v6 = ws.cell(r, 6).value
        v7 = ws.cell(r, 7).value
        if isinstance(v6, str) and "min/partit" in v6.lower() and isinstance(v7, str) and "punts/partit" in v7.lower():
            header_rows.append(r)
    return header_rows

def populate_template(resumen: pd.DataFrame, template_path: str, team_name: str=None, wins=None, losses=None):
    wb = load_workbook(template_path)
    ws = wb["Hoja1"]

    # --- Capçalera ---
    if team_name is not None:
        ws.cell(2, 2).value = team_name
    if wins is not None:
        ws.cell(3, 4).value = wins
    if losses is not None:
        ws.cell(3, 5).value = losses

    # --- Totals equip ---
    team_tot = team_totals_from_resumen(resumen)
    header_map = {
        6:  "PT",
        7:  "T2C",  8:  "T2I",  9:  "T2 (%)",
        10: "T3C",  11: "T3I", 12: "T3 (%)",
        13: "TCC",  14: "TCI", 15: "TC (%)",
        16: "TLC",  17: "TLI", 18: "TL (%)",
        19: "RO",   20: "RD",  21: "RT",
        22: "AS",   23: "BR",  24: "BP",
        25: "TapF", 26: "TapC", 27: "MT",
        28: "FC",   29: "FR",  30: "VA", 31: "Plays",
    }
    for c, key in header_map.items():
        if isinstance(team_tot, pd.Series) and key in team_tot.index and pd.notna(team_tot[key]):
            val = team_tot[key]
            if "%" in key:
                ws.cell(3, c).value = to_excel_percent(val)
            else:
                ws.cell(3, c).value = to_excel_number(val)
        else:
            ws.cell(3, c).value = None

    # --- Jugadores ---
    dfp = resumen.copy()
    if "Jugador" in dfp.columns:
        dfp = dfp[dfp["Jugador"].astype(str).apply(normalize_text) != "total"]
    if "MIN" in dfp.columns:
        dfp["MIN"] = pd.to_numeric(dfp["MIN"], errors="coerce")
        dfp = dfp.sort_values("MIN", ascending=False)
    dfp = dfp.reset_index(drop=True)

    card_map = {
        6: "MIN",
        7: "PT",
        8: "Partidos jugados",
        9: "T2C", 10:"T2I", 11:"T2 (%)", 12:"% US T2",
        13:"T3C", 14:"T3I", 15:"T3 (%)", 16:"% US T3",
        17:"TCC", 18:"TCI", 19:"TC (%)", 20:"% US TC",
        21:"TLC", 22:"TLI", 23:"TL (%)",
        24:"RO",  25:"RD",  26:"RT",   27:"AS",
        28:"VA",  29:"+/-", 30:"Plays",
    }

    header_rows = find_player_blocks(ws)
    for i, hr in enumerate(header_rows):
        if i >= len(dfp):
            break
        row_val = hr + 1
        player = dfp.iloc[i]

        # NOM i DORSAL
        ws.cell(hr, 5).value = player.get("Jugador", None)
        ws.cell(row_val, 5).value = player.get("Dorsal", None) if "Dorsal" in dfp.columns else None

        # Escriu cada mètrica (2 decimals màx.)
        for col_idx, key in card_map.items():
            if key is None:
                continue
            val = player.get(key, None) if key in dfp.columns else None
            if isinstance(val, (int, float, np.floating, np.integer)):
                if "%" in key:
                    ws.cell(row_val, col_idx).value = to_excel_percent(val)
                else:
                    ws.cell(row_val, col_idx).value = to_excel_number(val)
            else:
                ws.cell(row_val, col_idx).value = None if pd.isna(val) else val

    return wb

def save_overwrite_atomic(wb, target_path: str, backup=True):
    target = Path(target_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    if backup and target.exists():
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        bak = target.with_suffix(f".bak.{ts}.xlsx")
        shutil.copy2(target, bak)
    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx", dir=str(target.parent)) as tmpf:
        tmp_path = Path(tmpf.name)
    try:
        wb.save(tmp_path)
        os.replace(tmp_path, target)
    except Exception:
        if tmp_path.exists():
            tmp_path.unlink(missing_ok=True)
        raise

# ==========================
# ÚS: SOBREESCRIURE PLANTILLA
# ==========================
#template = r"C:\Users\aferrerr\OneDrive - Hospital Clínic de Barcelona\Escritorio\LF2 - 2025-56\LF2\scouting\j7-viladecans.xlsx"
template = r"C:\Users\aferrerr\OneDrive - Hospital Clínic de Barcelona\Escritorio\LF2 - 2025-56\LF2\scouting\j26-almeda.xlsx"

wb = populate_template(resumen, template, team_name=equip, wins=None, losses=None)
save_overwrite_atomic(wb, template, backup=True)



#resumen.to_excel(r"C:\Users\aferrerr\OneDrive - Hospital Clínic de Barcelona\Escritorio\LF2 - 2025-56\LF2\scouting\resumen-j23-islabonita.xlsx")





































